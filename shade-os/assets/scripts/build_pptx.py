from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
L="./assets/logos/"
A="assets/"
BLK=RGBColor(0x11,0x11,0x11); GRY=RGBColor(0x6E,0x6E,0x6E); WHT=RGBColor(0xFF,0xFF,0xFF)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
blank=prs.slide_layouts[6]
def bg(slide,img): 
    p=slide.shapes.add_picture(img,0,0,SW,SH); slide.shapes._spTree.remove(p._element); slide.shapes._spTree.insert(2,p._element); return p
def tb(slide,l,t,w,h,text,size,color=BLK,bold=False,align=PP_ALIGN.LEFT,spacing=0,font="Arial",anchor=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(l,t,w,h); tf=box.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text
    f=r.font; f.size=Pt(size); f.bold=bold; f.name=font; f.color.rgb=color
    return box
def logo(slide,white=False,right=True):
    img=L+("shade logo.png" if white else "shade logo 2.png")
    w=Inches(1.7); h=Inches(1.7*739/2922)
    l=SW-w-Inches(0.6) if right else Inches(0.6)
    slide.shapes.add_picture(img,l,Inches(0.45),w,h)

# --- Slide 1: title (dark) ---
s=prs.slides.add_slide(blank); bg(s,A+"bg_dark.png"); logo(s,white=True,right=False)
tb(s,Inches(0.9),Inches(3.0),Inches(11),Inches(0.6),"D E C O R A T I V E   L I G H T I N G   S E L E C T I O N",16,RGBColor(0xC9,0xC9,0xC9),False,PP_ALIGN.LEFT)
tb(s,Inches(0.85),Inches(3.4),Inches(11.5),Inches(1.2),"LUSAIL MARINA HEIGHTS",44,WHT,True,PP_ALIGN.LEFT)
tb(s,Inches(0.9),Inches(4.5),Inches(11),Inches(0.5),"MOCK-UP SELECTION    |    260608 - R00    |    JUNE 2026",14,RGBColor(0x9A,0x9A,0x9A),False,PP_ALIGN.LEFT)
# --- Slide 2: room plan ---
s=prs.slides.add_slide(blank); bg(s,A+"bg.png"); logo(s,white=False,right=True)
# vertical label
vb=tb(s,Inches(-2.4),Inches(3.2),Inches(5.6),Inches(0.5),"DECORATIVE LIGHT SELECTION",12,GRY,False,PP_ALIGN.CENTER,font="Arial")
vb.rotation=270
s.shapes.add_picture(A+"plan_ph.png",Inches(0.9),Inches(1.5),Inches(8.4),Inches(4.6))
tb(s,Inches(9.5),Inches(1.45),Inches(3),Inches(0.3),"KEY PLAN",11,GRY,True,PP_ALIGN.CENTER)
s.shapes.add_picture(A+"key_ph.png",Inches(9.5),Inches(1.8),Inches(2.9),Inches(3.1))
# room label bottom-left
tb(s,Inches(0.9),Inches(6.5),Inches(0.4),Inches(0.6),"/",30,BLK,True,PP_ALIGN.LEFT)
tb(s,Inches(1.3),Inches(6.55),Inches(6),Inches(0.6),"LIVING ROOM",22,BLK,False,PP_ALIGN.LEFT,spacing=2)
# --- Slide 3: fixture schedule grid ---
s=prs.slides.add_slide(blank); bg(s,A+"bg.png"); logo(s,white=False,right=True)
tb(s,Inches(0.9),Inches(0.5),Inches(8),Inches(0.6),"FIXTURE SELECTION",24,BLK,True,PP_ALIGN.LEFT)
tb(s,Inches(0.9),Inches(1.05),Inches(8),Inches(0.4),"LIVING ROOM  /  DECORATIVE",13,GRY,False,PP_ALIGN.LEFT)
fixtures=[("CW18A","Wall Sconce"),("TF-60","Pendant — Dining"),("TF-59","Pendant — Kitchen"),("3BST1","Table Lamp")]
x0,y0,gw,gh=Inches(0.9),Inches(1.7),Inches(3.0),Inches(2.4)
for i,(code,name) in enumerate(fixtures):
    cx=x0+ (Inches(3.05))*i
    s.shapes.add_picture(A+"fix_ph.png",cx,y0,Inches(2.7),Inches(2.7))
    tb(s,cx,y0+Inches(2.75),Inches(2.7),Inches(0.35),code,14,BLK,True,PP_ALIGN.LEFT)
    tb(s,cx,y0+Inches(3.1),Inches(2.7),Inches(0.35),name,12,GRY,False,PP_ALIGN.LEFT)
out="SHADE-Lighting-Selection-v1.pptx"; prs.save(out); print("saved",out)
