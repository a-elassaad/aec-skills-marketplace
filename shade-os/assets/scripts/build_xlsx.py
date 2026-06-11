import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PImage

BLK="FF111111"; WHT="FFFFFFFF"; SOFT="FFF4F2EE"; BAND="FFFAF9F7"; LINE="FFD8D5CF"
thin=Side(style="thin",color="FFD8D5CF")
border=Border(left=thin,right=thin,top=thin,bottom=thin)
def setw(ws,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

wb=openpyxl.Workbook()
ws=wb.active; ws.title="FF&E SCHEDULE"
ws.sheet_view.showGridLines=False
cols=["APARTMENT TYPE","ROOM","ITEM","CODE","CATEGORY","DIMENSION","QTY","REMARKS"]
setw(ws,[16,20,30,12,20,18,8,24])
# logo
logo_src="./assets/logos/shade logo 2.png"
im=PImage.open(logo_src); ratio=im.height/im.width
xi=XLImage(logo_src); xi.width=190; xi.height=int(190*ratio); ws.add_image(xi,"A1")
ws.row_dimensions[1].height=42; ws.row_dimensions[2].height=14
# title block (right side)
ws["F1"]="Project:"; ws["G1"]="LUSAIL MARINA HEIGHTS"
ws["F2"]="Title:";   ws["G2"]="BOQ — FF&E SCHEDULE"
ws["F3"]="Date:";    ws["G3"]="2026-06-10"
ws["F4"]="Ref:";     ws["G4"]="SG-M03 — R0"
for r in range(1,5):
    ws[f"F{r}"].font=Font(name="Arial",bold=True,size=9,color=BLK)
    ws[f"F{r}"].alignment=Alignment(horizontal="right")
    ws[f"G{r}"].font=Font(name="Arial",size=9,color="FF333333")
HEADROW=6
for c,name in enumerate(cols,1):
    cell=ws.cell(row=HEADROW,column=c,value=name)
    cell.fill=PatternFill("solid",fgColor=BLK); cell.font=Font(name="Arial",bold=True,size=9,color=WHT)
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=border
ws.row_dimensions[HEADROW].height=26
data=[
 ["1BHK","Living / Dining","Dining Table 4 pax","1B.DT4","FURNITURE","Ø 100 cm",1,""],
 ["","","Dining Chairs","1B.DCH1","FURNITURE","36 x 45 cm",4,""],
 ["","","L Sofa","1B.SF","FURNITURE","380 x 200 cm",1,""],
 ["","","Coffee Table","1B.CT1","FURNITURE","Ø 100 cm",1,""],
 ["","","Arm Chair 1","1B.ACH1","FURNITURE","75 x 50 cm",2,""],
 ["","","TV Unit","1B.TVU1","FURNITURE","refer to dwgs",1,"incl. drawers"],
 ["3BHK","Living / Dining","Dining Table 6 pax","3B.DT6","FURNITURE","80 x 175 cm",1,""],
 ["","","Dining Chairs","3B.DCH1","FURNITURE","36 x 45 cm",6,""],
 ["","","L Sofa","3B.SF","FURNITURE","80 x 312 cm",1,""],
 ["","","Side Table","3B.ST","FURNITURE","45 x 45 cm",2,""],
 ["","","Arm Chair 1","3B.ACH1","FURNITURE","75 x 50 cm",2,""],
 ["","Kitchen","Built-in Refrigerator","3B.APP1","KITCHEN APPLIANCES","TBD",1,""],
 ["","Kitchen","Built-in Oven","3B.APP2","KITCHEN APPLIANCES","60 x 60 cm",1,""],
 ["","Master Bath","Wash Basin","3B.SAN1","SANITARY WARES","TBD",2,""],
 ["","Master Bath","Wash Mixer","3B.SAN2","SANITARY WARES","TBD",2,""],
 ["","Living / Dining","Pendant Light","3B.LT1","LIGHTING","TBD",1,"decorative"],
]
r=HEADROW+1
for i,row in enumerate(data):
    for c,val in enumerate(row,1):
        cell=ws.cell(row=r,column=c,value=val); cell.border=border
        cell.font=Font(name="Arial",size=9,color=BLK)
        cell.alignment=Alignment(horizontal="center" if c in(4,7) else "left",vertical="center",wrap_text=True)
        if i%2==1: cell.fill=PatternFill("solid",fgColor=BAND)
    r+=1
# total row
tot=r
ws.cell(row=tot,column=6,value="TOTAL").font=Font(name="Arial",bold=True,size=9)
ws.cell(row=tot,column=6).alignment=Alignment(horizontal="right")
tc=ws.cell(row=tot,column=7,value=f"=SUM(G{HEADROW+1}:G{r-1})")
tc.font=Font(name="Arial",bold=True,size=9,color=WHT); tc.fill=PatternFill("solid",fgColor=BLK)
tc.alignment=Alignment(horizontal="center"); tc.border=border
ws.freeze_panes=f"A{HEADROW+1}"
ws.print_title_rows=f"{HEADROW}:{HEADROW}"
# data validation for CATEGORY
dv=DataValidation(type="list",formula1='"FURNITURE,KITCHEN APPLIANCES,SANITARY WARES,LIGHTING,COMMON AREA FURNITURE,OUTDOOR FURNITURE,ACCESSORIES"',allow_blank=True)
ws.add_data_validation(dv); dv.add(f"E{HEADROW+1}:E{r+50}")

# SUMMARY sheet
s=wb.create_sheet("SUMMARY"); s.sheet_view.showGridLines=False
setw(s,[28,14])
s["A1"]="FF&E SUMMARY — BY CATEGORY"; s["A1"].font=Font(name="Arial",bold=True,size=12,color=BLK)
cats=["FURNITURE","KITCHEN APPLIANCES","SANITARY WARES","LIGHTING","COMMON AREA FURNITURE","OUTDOOR FURNITURE","ACCESSORIES"]
s["A3"]="CATEGORY"; s["B3"]="TOTAL QTY"
for cc in ("A3","B3"):
    s[cc].fill=PatternFill("solid",fgColor=BLK); s[cc].font=Font(name="Arial",bold=True,size=9,color=WHT)
    s[cc].alignment=Alignment(horizontal="center"); s[cc].border=border
rr=4
for cat in cats:
    s.cell(row=rr,column=1,value=cat).font=Font(name="Arial",size=9)
    s.cell(row=rr,column=1).border=border
    f=f"=SUMIF('FF&E SCHEDULE'!E:E,A{rr},'FF&E SCHEDULE'!G:G)"
    cell=s.cell(row=rr,column=2,value=f); cell.alignment=Alignment(horizontal="center"); cell.border=border; cell.font=Font(name="Arial",size=9)
    rr+=1
s.cell(row=rr,column=1,value="TOTAL").font=Font(name="Arial",bold=True,size=9)
t=s.cell(row=rr,column=2,value=f"=SUM(B4:B{rr-1})"); t.font=Font(name="Arial",bold=True,size=9,color=WHT); t.fill=PatternFill("solid",fgColor=BLK); t.alignment=Alignment(horizontal="center"); t.border=border

out="SHADE-FFE-BOQ-v1.xlsx"; wb.save(out); print("saved",out)
